import openpyxl

def inspect_fast():
    print("=== RAPID READ ONLY INSPECTION ===")
    wb = openpyxl.load_workbook('docs/ATALA IA MODEL (1).xlsm', read_only=True)
    sheet = wb['EURGBPUSD']
    
    # Encontrar las primeras filas con datos no nulos
    found_rows = 0
    for r in range(1, 1000):
        # Leer algunas columnas
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 40)]
        # Si la fila tiene al menos algún valor no nulo
        if any(v is not None for v in row_vals):
            found_rows += 1
            # Imprimir la fila
            non_nulls = {openpyxl.utils.get_column_letter(c+1): v for c, v in enumerate(row_vals) if v is not None}
            print(f"Row {r}: {non_nulls}")
            if found_rows > 30:
                break

if __name__ == '__main__':
    inspect_fast()
